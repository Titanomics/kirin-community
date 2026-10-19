"""Real engine + durable snapshot adapter, with synthetic clinics and workbooks."""
import base64
import io
import json
import unittest
from openpyxl import Workbook
from analytics_runtime.backend import CloudError
from analytics_runtime.runtime import dispatch

FIRST='11111111-1111-4111-8111-111111111111'
SECOND='22222222-2222-4222-8222-222222222222'
USER='aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa'


def workbook(count=6,amount=100):
    book=Workbook();sheet=book.active
    sheet.append(['접수시각','차트번호','신구','최초내원','총진료비','카드수납','현금수납','기타(온라인)','주소'])
    for i in range(count): sheet.append(['2024-02-02','SYNTHETIC-'+str(i),'신환','2024-02-02',amount,amount,0,0,'서울특별시 강남구 역삼동'])
    out=io.BytesIO();book.save(out);book.close();return out.getvalue()


def channel_workbook(count=6):
    book=Workbook();sheet=book.active
    sheet.append(['내원경로','내원환자수','신환수','총내원횟수','총진료비'])
    sheet.append(['합계',count,count,count,count*100])
    detail=book.create_sheet('환자');detail.append(['차트번호','주소'])
    for i in range(count):detail.append(['SYNTHETIC-'+str(i),'서울특별시 강남구 역삼동'])
    out=io.BytesIO();book.save(out);book.close();return out.getvalue()


class MemoryBackend:
    def __init__(self):
        self.states={cid:({'revision':0},None) for cid in (FIRST,SECOND)}
        self.role='admin';self.allowed={FIRST,SECOND};self.conflict=False;self.objects={}
    def authorize(self,clinic_id,write=False):
        if clinic_id not in self.allowed or (write and self.role=='viewer'): raise CloudError('Denied',403)
        return {'id':USER,'username':USER,'name':'TEST OPERATOR','role':self.role,'version':1,'clinic_ids':list(self.allowed)}, {'id':clinic_id,'name':'SYNTHETIC A' if clinic_id==FIRST else 'SYNTHETIC B'}
    def read_snapshot(self,clinic_id): return self.states[clinic_id]
    def commit(self,clinic_id,revision,data):
        if self.conflict or self.states[clinic_id][0]['revision']!=revision: raise CloudError('Conflict',409)
        self.states[clinic_id]=({'revision':revision+1},data);return revision+1
    def download(self,clinic_id,path,kind,limit):
        if not path.startswith(clinic_id+'/'+kind+'/'): raise CloudError('Denied',403)
        return self.objects[path]


class EngineFlowTests(unittest.TestCase):
    def setUp(self): self.backend=MemoryBackend()
    def call(self,path,payload=None,clinic=FIRST):
        result=dispatch(self.backend,{'clinic':clinic,'method':'GET' if payload is None else 'POST','path':path,'payload':payload})
        body=json.loads(result[1]) if result[2].startswith('application/json') else result[1].decode('utf8')
        return result[0],body
    def import_data(self,clinic=FIRST,count=6):
        status,preview=self.call('/api/imports/preview',{'filename':'synthetic.xlsx','year':2024,'kind':'receipts','data':base64.b64encode(workbook(count)).decode()},clinic)
        self.assertEqual(status,200,preview)
        status,result=self.call('/api/imports/commit',{'id':preview['id'],'revision':preview['revision'],'confirmed':True},clinic)
        self.assertEqual(status,200,result)
        self.assertEqual(result['status'],'committed')
    def test_cold_start_preview_commit_dashboard_and_frozen_presentation(self):
        self.import_data()
        status,dashboard=self.call('/api/dashboard');self.assertEqual(status,200,dashboard)
        self.assertEqual(dashboard['current']['patients'],6)
        self.assertEqual(dashboard['current']['collection_recorded'],600)
        self.assertNotIn('SYNTHETIC-0',json.dumps(dashboard))
        status,result=self.call('/api/workspace/reports/generate',{'month':'2024-02','revision':0,'notes':{}})
        self.assertEqual(status,200,result);rid=result['entity']['id']
        status,html=self.call(f'/api/workspace/reports/{rid}/presentation');self.assertEqual(status,200)
        self.assertIn('SYNTHETIC A',html);self.assertNotIn('SYNTHETIC B',html);self.assertNotIn('SYNTHETIC-0',html)
        self.import_data(count=9)
        self.assertEqual(self.call(f'/api/workspace/reports/{rid}/presentation')[1],html)
        self.assertEqual(self.call(f'/api/workspace/reports/{rid}/presentation',clinic=SECOND)[0],404)
    def test_two_populated_clinics_remain_independent(self):
        self.import_data(FIRST,6);self.import_data(SECOND,9)
        self.assertEqual(self.call('/api/dashboard',clinic=FIRST)[1]['current']['patients'],6)
        self.assertEqual(self.call('/api/dashboard',clinic=SECOND)[1]['current']['patients'],9)
        self.backend.allowed={FIRST};self.assertEqual(self.call('/api/dashboard',clinic=SECOND)[0],403)
        self.assertEqual(self.call('/api/dashboard?clinic='+SECOND)[0],403)
    def test_viewer_cannot_import_and_conflict_does_not_overwrite(self):
        self.import_data();before=self.backend.states[FIRST]
        self.backend.role='viewer'
        self.assertEqual(self.call('/api/imports/rollback',{'revision':1,'confirmed':True})[0],403)
        self.backend.role='editor';self.backend.conflict=True
        self.assertEqual(self.call('/api/imports/rollback',{'revision':1,'confirmed':True})[0],409)
        self.assertEqual(self.backend.states[FIRST],before)
    def test_cross_clinic_upload_and_account_substitution_are_rejected(self):
        result=self.call('/api/imports/inspect',{'filename':'x.xlsx','upload_path':SECOND+'/uploads/test.xlsx'})
        self.assertEqual(result[0],403)
        result=dispatch(self.backend,{'clinic':FIRST,'method':'GET','path':'/api/imports','identity':'other:1'})
        self.assertEqual(result[0],412)
    def test_corrupted_snapshot_is_not_reset(self):
        self.backend.states[FIRST]=({'revision':4},b'corrupt')
        self.assertEqual(self.call('/api/imports')[0],503)
        self.assertEqual(self.backend.states[FIRST][0]['revision'],4)
    def test_local_accounts_and_filesystem_operations_not_exposed(self):
        for path in ('/api/auth/setup','/api/agency/clinics','/api/recovery/restore','/../secrets'):
            self.assertGreaterEqual(self.call(path,{})[0],400)
        status,html=self.call('/index.html');self.assertEqual(status,200)
        self.assertIn('/api/hospital-analytics/runtime/'+FIRST+'/app.js',html)
        self.assertNotIn('디데이치과',html)


if __name__=='__main__':unittest.main()
