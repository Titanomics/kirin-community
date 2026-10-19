import io,json,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook
from analytics_runtime.snapshots import pack
from analytics_runtime.runtime import dispatch
from test_engine_flow import MemoryBackend,FIRST,SECOND

def aggregate_book(total=10,rank_count=7):
    book=Workbook();sheet=book.active;sheet.title='아파트별 집계'
    sheet.append(['검증병원 | 아파트별 신규환자 주소 집계']);sheet.append([]);sheet.append([])
    sheet.append(['번호','아파트명','도로명주소','환자 기록(건)','2개 단지 내 비중'])
    sheet.append([1,'검증단지 A','서울특별시 강남구 역삼동',7,.7])
    sheet.append([2,'검증단지 B','서울특별시 강남구 삼성동',3,.3])
    sheet.append([None,None,'2개 단지 합계',total,1])
    rank=book.create_sheet('환자수 순위');rank.append(['순위','아파트명','환자 기록(건)','비중'])
    rank.append([1,'검증단지 A',rank_count,.7]);rank.append([2,'검증단지 B',3,.3])
    out=io.BytesIO();book.save(out);return out.getvalue()

class AddressSummaryTests(unittest.TestCase):
    def parse(self,**kwargs):
        from analytics_runtime.address_summary import parse_workbook
        return parse_workbook(aggregate_book(**kwargs),period='2026-09',source_name='aggregate.xlsx')
    def test_validates_both_sheets_and_preserves_record_definition(self):
        report=self.parse()
        self.assertEqual(report['total_records'],10)
        self.assertEqual(len(report['rows']),2)
        self.assertEqual(report['period'],'2026-09')
        self.assertEqual(report['measure'],'신규환자 기록')
        self.assertNotIn('unique_patients',report)
    def test_rejects_total_or_ranking_disagreement(self):
        from analytics_runtime.address_summary import SummaryError
        for values in ({'total':11},{'rank_count':8}):
            with self.subTest(values=values),self.assertRaises(SummaryError):self.parse(**values)
    def test_sparse_worksheet_dimensions_are_bounded_before_reading(self):
        from analytics_runtime.address_summary import parse_workbook,SummaryError
        from unittest.mock import patch
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        for row,column in ((1048576,1),(1,16384)):
            book=Workbook();book.active.cell(row,column,'sparse');out=io.BytesIO();book.save(out);book.close()
            with patch.object(ReadOnlyWorksheet,'iter_rows',side_effect=AssertionError('Must reject before iterating')):
                with self.assertRaises(SummaryError):parse_workbook(out.getvalue(),period='2026-09',source_name='sparse.xlsx')
    def test_rejects_corrupt_stored_aggregates_without_resetting_snapshot(self):
        for changed in ({'total_records':11},{'clinic_id':SECOND},{'period':'2026-10'},{'rows':[{'name':'bad','address':'x','records':-1}]}):
            report={**self.parse(),'clinic_id':FIRST,**changed};backend=MemoryBackend()
            with tempfile.TemporaryDirectory() as folder:
                root=Path(folder);(root/'address-reports').mkdir()
                (root/'address-reports'/'2026-09.json').write_text(json.dumps(report),encoding='utf8')
                backend.commit(FIRST,0,pack(root))
            before=backend.states[FIRST]
            response=dispatch(backend,{'clinic':FIRST,'method':'GET','path':'/index.html'})
            self.assertEqual(response[0],503,changed);self.assertEqual(backend.states[FIRST],before)
    def test_multiple_periods_remain_independent(self):
        backend=MemoryBackend()
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);(root/'address-reports').mkdir()
            for period,total,rank in [('2026-08',10,7),('2026-09',11,8)]:
                report=self.parse();report.update(clinic_id=FIRST,period=period,total_records=total)
                report['rows'][0]['records']=rank
                (root/'address-reports'/f'{period}.json').write_text(json.dumps(report),encoding='utf8')
            backend.commit(FIRST,0,pack(root))
        def get(path):return dispatch(backend,{'clinic':FIRST,'method':'GET','path':path})
        self.assertEqual([r['total_records'] for r in json.loads(get('/api/address-reports')[1])['reports']],[10,11])
        self.assertIn(b'2026-08',get('/index.html?month=2026-08')[1])
        self.assertEqual(get('/index.html?month=2026-07')[0],404)
        self.assertEqual(get('/index.html?clinic='+SECOND)[0],403)
    def test_report_is_standalone_escaped_and_not_fabricated_revenue(self):
        from analytics_runtime.address_summary import render_report
        html=render_report(self.parse(),{'id':FIRST,'name':'<script>alert(1)</script>'})
        self.assertIn('&lt;script&gt;',html)
        self.assertIn('10',html)
        self.assertNotIn('fetch(',html)
        self.assertNotIn('src="http',html)
        self.assertNotIn('고유환자 10',html)
        self.assertIn('중복 환자 여부',html)
    def test_cloud_aggregate_persists_and_cannot_cross_clinic(self):
        report=self.parse();report['clinic_id']=FIRST
        backend=MemoryBackend()
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);(root/'address-reports').mkdir()
            (root/'address-reports'/'2026-09.json').write_text(json.dumps(report),encoding='utf8')
            backend.commit(FIRST,0,pack(root))
        def get(path,clinic=FIRST):return dispatch(backend,{'clinic':clinic,'method':'GET','path':path})
        self.assertEqual(get('/index.html')[0],200)
        self.assertIn('신규환자 기록'.encode(),get('/index.html')[1])
        self.assertEqual(get('/api/address-reports/2026-09/html')[0],200)
        self.assertEqual(get('/api/address-reports/2026-09/html',SECOND)[0],404)
        self.assertNotIn(b'data-page="quality"',get('/index.html')[1])
        backend.allowed={SECOND}
        self.assertEqual(get('/api/address-reports/2026-09/html')[0],403)
    def test_legacy_workspace_has_no_data_update_or_mapping_ui(self):
        backend=MemoryBackend()
        def get(path):return dispatch(backend,{'clinic':FIRST,'method':'GET','path':path})
        self.assertNotIn(b'data-page="quality"',get('/index.html')[1])
        self.assertNotIn('열 연결'.encode(),get('/imports.js')[1])
        backend.role='editor'
        result=dispatch(backend,{'clinic':FIRST,'method':'POST','path':'/api/imports/preview','payload':{}})
        self.assertEqual(result[0],403)

if __name__=='__main__':unittest.main()
