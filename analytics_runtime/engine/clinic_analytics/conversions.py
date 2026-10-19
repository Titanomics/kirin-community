"""Local, opaque-case event imports. Only aggregate reports leave this store."""
import csv
import hashlib
import io
import json
import sqlite3
import threading
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime,date,timedelta
from pathlib import Path
import openpyxl
from .loader import DataError,KST
from .marketing import now,text,period,ACTOR
from .knowledge import day,today

HEADERS={'이벤트ID':'event_id','연결키':'subject_id','단계':'stage','발생일시':'occurred_at','캠페인ID':'campaign_id','금액':'amount','원거래ID':'refund_of'}
STAGES={'inquiry':'문의','booking':'예약','attendance':'내원','consultation':'상담','treatment':'치료','payment':'수납','refund':'환불','cancellation':'예약취소','no_show':'미내원'}
INIT_LOCK=threading.Lock()
MAX_BYTES=5*1024*1024
MAX_ROWS=25000


def opaque(value):
    if not isinstance(value,str):raise ValueError
    parsed=uuid.UUID(value.strip())
    if parsed.version!=4 or parsed.variant!=uuid.RFC_4122:raise ValueError
    return parsed.hex


def parse_upload(filename,data):
    if not isinstance(data,bytes) or not 0<len(data)<=MAX_BYTES:raise DataError('연결자료는 5MB 이하의 CSV 또는 XLSX 파일로 선택해 주세요.')
    ext=Path(str(filename)).suffix.lower();book=None
    try:
        if ext=='.csv':
            try:decoded=data.decode('utf-8-sig')
            except UnicodeDecodeError:decoded=data.decode('cp949')
            rows=csv.reader(io.StringIO(decoded))
        elif ext=='.xlsx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries=archive.infolist()
                if len(entries)>1000 or sum(x.file_size for x in entries)>50*1024*1024:raise ValueError
                for entry in entries:
                    if entry.flag_bits&1 or entry.filename.endswith('vbaProject.bin'):raise ValueError
                    if entry.filename.endswith(('.xml','.rels')):
                        content=archive.read(entry).upper()
                        if b'<!DOCTYPE' in content or b'<!ENTITY' in content:raise ValueError
            book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            if len(book.worksheets)!=1:raise ValueError
            rows=book.worksheets[0].iter_rows(values_only=True)
        else:raise ValueError
        headers=[str(v or '').strip() for v in next(rows,[])]
        if len(headers)!=len(HEADERS) or set(headers)!=set(HEADERS):raise DataError('연결자료 양식의 7개 열을 사용해 주세요. 이름·전화번호·차트번호 열은 받지 않습니다.')
        result=[];seen=set();reverse={v:k for k,v in STAGES.items()}
        for line,values in enumerate(rows,2):
            if line>MAX_ROWS+1:raise DataError('한 번에 25,000개 이하의 행을 올려 주세요.')
            if not any(v is not None and str(v).strip() for v in values):continue
            if len(result)>=MAX_ROWS:raise DataError('한 번에 25,000개 이하의 이벤트를 올려 주세요.')
            try:
                if len(values)!=len(headers):raise ValueError
                row={HEADERS[h]:v for h,v in zip(headers,values)}
                event_id=opaque(row['event_id']);subject=opaque(row['subject_id']);campaign=opaque(row['campaign_id'])
                if event_id in seen:raise ValueError
                stage=str(row['stage']).strip();stage=reverse.get(stage,stage)
                if stage not in STAGES:raise ValueError
                value=row['occurred_at'];stamp=value if isinstance(value,datetime) else datetime.fromisoformat(str(value).strip())
                if stamp.tzinfo is None:stamp=stamp.replace(tzinfo=KST)
                stamp=stamp.astimezone(KST)
                if stamp>datetime.now(KST) or stamp.year<2000:raise ValueError
                amount=row['amount'];amount=0 if amount is None or str(amount).strip()=='' else amount
                if isinstance(amount,bool) or not str(amount).isdigit():raise ValueError
                amount=int(amount)
                if not 0<=amount<=10**12 or (stage in ('payment','refund') and amount==0) or (stage not in ('payment','refund') and amount):raise ValueError
                refund=opaque(row['refund_of']) if row['refund_of'] else None
                if (stage=='refund')!=bool(refund):raise ValueError
                result.append({'event_id':event_id,'subject_id':subject,'campaign_id':campaign,'stage':stage,'occurred_at':stamp.isoformat(),'amount':amount,'refund_of':refund});seen.add(event_id)
            except (ValueError,TypeError,AttributeError):raise DataError(f'{line}행의 연결키·단계·시각·금액을 확인해 주세요. ID는 무작위 UUID v4 형식이어야 합니다.') from None
        if not result:raise DataError('입력된 이벤트가 없습니다. 양식에 연결자료를 입력해 주세요.')
        return result
    except DataError:raise
    except Exception:raise DataError('파일을 읽지 못했습니다. UTF-8 CSV 또는 암호·수식 없는 단일 시트 XLSX를 사용해 주세요.') from None
    finally:
        if book:book.close()


class ConversionStore:
    def __init__(self,path,clinic_id):self.path=Path(path);self.clinic_id=clinic_id;self.marker=self.path.with_suffix(self.path.suffix+'.initialized')

    @contextmanager
    def connection(self,write=False):
        conn=None
        try:
            self.path.parent.mkdir(parents=True,exist_ok=True)
            with INIT_LOCK:
                exists=self.path.exists()
                if not exists and self.marker.exists():raise DataError('연결자료 저장소가 없습니다. 빈 자료로 초기화하지 않습니다.')
                if not exists:self.marker.write_text('initialized\n',encoding='ascii')
                conn=sqlite3.connect(self.path,timeout=25,isolation_level=None);conn.row_factory=sqlite3.Row
                if not exists:
                    conn.executescript('''CREATE TABLE meta(id INTEGER PRIMARY KEY CHECK(id=1),clinic_id TEXT,revision INTEGER);
                    CREATE TABLE batches(id TEXT PRIMARY KEY,sha256 TEXT,format TEXT,observed_through TEXT,created_at TEXT,inserted INTEGER,skipped INTEGER,voided INTEGER,void_reason TEXT,actor TEXT);
                    CREATE TABLE events(batch_id TEXT,event_id TEXT,subject_id TEXT,campaign_id TEXT,stage TEXT,occurred_at TEXT,amount INTEGER,refund_of TEXT,PRIMARY KEY(batch_id,event_id));
                    CREATE TABLE coverage(batch_id TEXT,subject_id TEXT,PRIMARY KEY(batch_id,subject_id));
                    CREATE TABLE audit(revision INTEGER PRIMARY KEY,action TEXT,batch_id TEXT,at TEXT,actor TEXT,reason TEXT);''')
                    conn.execute('INSERT INTO meta VALUES(1,?,0)',(self.clinic_id,))
                else:
                    conn.execute('SELECT id,clinic_id,revision FROM meta LIMIT 0');conn.execute('SELECT id,voided,sha256 FROM batches LIMIT 0')
                    conn.execute('SELECT batch_id,event_id,subject_id,campaign_id,stage,occurred_at,amount,refund_of FROM events LIMIT 0')
                    conn.execute('SELECT batch_id,subject_id FROM coverage LIMIT 0');conn.execute('SELECT revision,action,batch_id,at,actor,reason FROM audit LIMIT 0')
                    if not self.marker.exists():self.marker.write_text('initialized\n',encoding='ascii')
            conn.execute('BEGIN IMMEDIATE' if write else 'BEGIN')
            meta=conn.execute('SELECT * FROM meta WHERE id=1').fetchone()
            if not meta or meta['clinic_id']!=self.clinic_id:raise DataError('연결자료 병원 정보가 일치하지 않습니다.')
            yield conn,meta['revision'];conn.commit()
        except (sqlite3.Error,OSError,ValueError,TypeError,KeyError):
            if conn:conn.rollback()
            raise DataError('연결자료 저장소를 확인하지 못했습니다.') from None
        except Exception:
            if conn:conn.rollback()
            raise
        finally:
            if conn:conn.close()

    def _events(self,conn):return [dict(r) for r in conn.execute('SELECT e.event_id,e.subject_id,e.campaign_id,e.stage,e.occurred_at,e.amount,e.refund_of FROM events e JOIN batches b ON b.id=e.batch_id WHERE b.voided=0')]

    def _validate(self,conn,rows,through,workspace):
        through=day(through,'자료 확인 기준일',True)
        if through>today() or any(r['occurred_at'][:10]>through for r in rows):raise DataError('자료 기준일은 이벤트 발생일 이후이면서 오늘 이전이어야 합니다.')
        if workspace.get('clinic_id')!=self.clinic_id:raise DataError('병원 작업공간이 일치하지 않습니다.')
        campaigns={x['id'] for x in workspace['campaigns']}
        if any(r['campaign_id'] not in campaigns for r in rows):raise DataError('이 병원에 등록된 캠페인 ID를 사용해 주세요.')
        existing={x['event_id']:x for x in self._events(conn)};new=[];skipped=0
        for row in rows:
            old=existing.get(row['event_id'])
            if old:
                if row!=old:raise DataError('같은 이벤트 ID의 내용이 다릅니다. 기존 업로드를 정정한 뒤 다시 올려 주세요.')
                skipped+=1
            else:new.append(row)
        combined={**existing,**{x['event_id']:x for x in new}};case_campaign={};refunds={}
        for row in combined.values():
            previous=case_campaign.setdefault(row['subject_id'],row['campaign_id'])
            if previous!=row['campaign_id']:raise DataError('한 연결키는 하나의 캠페인에만 연결합니다. 중복 귀속을 확인해 주세요.')
            if row['stage']=='refund':
                original=combined.get(row['refund_of'])
                if not original or original['stage']!='payment' or original['subject_id']!=row['subject_id'] or original['campaign_id']!=row['campaign_id'] or original['occurred_at']>row['occurred_at']:raise DataError('환불은 같은 연결건의 선행 수납 이벤트를 지정해 주세요.')
                refunds[original['event_id']]=refunds.get(original['event_id'],0)+row['amount']
                if refunds[original['event_id']]>original['amount']:raise DataError('누적 환불액이 원 수납액을 초과합니다.')
        if len(combined)>200000:raise DataError('이 작업공간의 연결자료 한도를 초과했습니다.')
        return new,skipped,through

    def preview(self,filename,data,through,workspace):
        rows=parse_upload(filename,data);sha=hashlib.sha256(data).hexdigest()
        with self.connection() as (conn,revision):
            new,skipped,through=self._validate(conn,rows,through,workspace)
            return {'sha256':sha,'revision':revision,'workspace_revision':workspace['revision'],'rows':len(rows),'inserted':len(new),'skipped':skipped,'cases':len({x['subject_id'] for x in rows}),
                'observed_through':through,'first_event':min(x['occurred_at'][:10] for x in rows),'last_event':max(x['occurred_at'][:10] for x in rows),'stages':{label:sum(x['stage']==stage for x in rows) for stage,label in STAGES.items()},'checked_at':now()}

    def _audit(self,conn,revision,action,batch,reason=''):
        principal=ACTOR.get();actor=json.dumps({k:principal[k] for k in ('id','username','name')} if principal else None,ensure_ascii=False)
        conn.execute('UPDATE meta SET revision=? WHERE id=1',(revision+1,))
        conn.execute('INSERT INTO audit VALUES(?,?,?,?,?,?)',(revision+1,action,batch,now(),actor,reason))
        return actor

    def commit(self,filename,data,through,workspace,revision,sha,confirmed):
        if confirmed is not True:raise DataError('미리보기와 연결자료의 출처·기준일을 확인해 주세요.')
        if not isinstance(sha,str) or hashlib.sha256(data).hexdigest()!=sha:raise DataError('미리보기 이후 파일이 바뀌었습니다. 다시 확인해 주세요.')
        rows=parse_upload(filename,data)
        with self.connection(True) as (conn,current):
            if type(revision)!=int or revision!=current:raise DataError('다른 화면에서 연결자료가 바뀌었습니다. 미리보기를 다시 확인해 주세요.')
            new,skipped,through=self._validate(conn,rows,through,workspace)
            if conn.execute('SELECT 1 FROM batches WHERE sha256=? AND observed_through=? AND voided=0',(sha,through)).fetchone():return {'revision':current,'duplicate':True,'inserted':0,'skipped':len(rows)}
            if conn.execute('SELECT count(*) FROM batches').fetchone()[0]>=5000:raise DataError('업로드 이력 한도에 도달했습니다.')
            identifier=uuid.uuid4().hex;actor=self._audit(conn,current,'import',identifier)
            conn.execute('INSERT INTO batches VALUES(?,?,?,?,?,?,?,0,?,?)',(identifier,sha,Path(filename).suffix.lower(),through,now(),len(new),skipped,'',actor))
            conn.executemany('INSERT INTO events VALUES(?,?,?,?,?,?,?,?)',[(identifier,*(r[k] for k in ('event_id','subject_id','campaign_id','stage','occurred_at','amount','refund_of'))) for r in new])
            conn.executemany('INSERT INTO coverage VALUES(?,?)',[(identifier,s) for s in {x['subject_id'] for x in rows}])
            return {'revision':current+1,'duplicate':False,'inserted':len(new),'skipped':skipped,'batch_id':identifier}

    def rollback(self,revision,reason,confirmed):
        reason=text(reason,'되돌리기 사유',300,True)
        if confirmed is not True:raise DataError('최근 업로드를 되돌릴지 확인해 주세요.')
        with self.connection(True) as (conn,current):
            if type(revision)!=int or revision!=current:raise DataError('연결자료가 바뀌었습니다. 다시 확인해 주세요.')
            batch=conn.execute("SELECT b.id FROM batches b JOIN audit a ON a.batch_id=b.id AND a.action='import' WHERE b.voided=0 ORDER BY a.revision DESC LIMIT 1").fetchone()
            if not batch:raise DataError('되돌릴 업로드가 없습니다.')
            conn.execute('UPDATE batches SET voided=1,void_reason=? WHERE id=?',(reason,batch['id']));self._audit(conn,current,'rollback',batch['id'],reason)
            return {'revision':current+1,'batch_id':batch['id']}

    def status(self):
        with self.connection() as (conn,revision):
            batches=[dict(r) for r in conn.execute("SELECT b.id,b.format,b.observed_through,b.created_at,b.inserted,b.skipped,b.voided,b.void_reason FROM batches b JOIN audit a ON a.batch_id=b.id AND a.action='import' ORDER BY a.revision DESC")]
            return {'revision':revision,'checked_at':now(),'batches':batches,'active_events':conn.execute('SELECT count(*) FROM events e JOIN batches b ON b.id=e.batch_id WHERE b.voided=0').fetchone()[0]}

    def report(self,start,end,window,workspace):
        if workspace.get('clinic_id')!=self.clinic_id:raise DataError('병원 작업공간이 일치하지 않습니다.')
        start,end=period(start,end)
        if type(window)!=int or window not in (30,60,90):raise DataError('관측 기간은 30·60·90일 중 선택해 주세요.')
        with self.connection() as (conn,revision):
            events=self._events(conn)
            coverage={r[0]:r[1] for r in conn.execute('SELECT c.subject_id,max(b.observed_through) FROM coverage c JOIN batches b ON b.id=c.batch_id WHERE b.voided=0 GROUP BY c.subject_id')}
        return aggregate(events,coverage,start,end,window,workspace,revision)


def aggregate(events,coverage,start,end,window,workspace,revision):
    groups={}
    for event in events:groups.setdefault(event['subject_id'],[]).append(event)
    rows={};orphan=complete=incomplete=0;cutoffs=[];anchors=0;observed_dates=[];window_ends=[]
    def blank():return {**{k:0 for k in STAGES},'collections':0,'refunds':0,'net_collections':0,'net_spend':0,'expense_entries':0,'complete_cases':0,'incomplete_cases':0}
    totals=blank()
    for subject,case in groups.items():
        inquiries=[x for x in case if x['stage']=='inquiry']
        if not inquiries:
            if any(start<=x['occurred_at'][:10]<=end for x in case):orphan+=1
            continue
        anchor=min(inquiries,key=lambda x:x['occurred_at']);anchor_day=anchor['occurred_at'][:10]
        if not start<=anchor_day<=end:continue
        anchors+=1;cutoff=coverage[subject];cutoffs.append(cutoff)
        last=(date.fromisoformat(anchor_day)+timedelta(days=window-1)).isoformat()
        window_ends.append(last)
        mature=cutoff>=last and last<today();complete+=int(mature);incomplete+=int(not mature)
        row=rows.setdefault(anchor['campaign_id'],blank());row['complete_cases']+=int(mature);row['incomplete_cases']+=int(not mature)
        observed=[x for x in case if anchor['occurred_at']<=x['occurred_at'] and x['occurred_at'][:10]<=min(last,cutoff)]
        observed_dates.extend(x['occurred_at'][:10] for x in observed)
        for stage in {x['stage'] for x in observed}:totals[stage]+=1;row[stage]+=1
        for event in observed:
            if event['stage'] in ('payment','refund'):
                key='collections' if event['stage']=='payment' else 'refunds';totals[key]+=event['amount'];row[key]+=event['amount']
    for expense in workspace['expenses']:
        if expense.get('voided') or not start<=expense['date']<=end:continue
        row=rows.setdefault(expense['campaign_id'],blank());value=expense['amount']*(-1 if expense['direction']=='refund' else 1)
        row['net_spend']+=value;totals['net_spend']+=value
        row['expense_entries']+=1;totals['expense_entries']+=1
    has_data=bool(anchors);campaigns={x['id']:x for x in workspace['campaigns']}
    for key,row in rows.items():
        row.update(campaign_id=key,title=campaigns.get(key,{}).get('title','캠페인 확인 필요'))
        row['net_collections']=row['collections']-row['refunds'] if row['inquiry'] else None
        row['attendance_rate']=round(row['attendance']/row['inquiry']*100,2) if row['inquiry'] else None
        row['cost_per_linked_attendance']=round(row['net_spend']/row['attendance']) if row['attendance'] and row['expense_entries'] and not row['incomplete_cases'] and row['net_spend']>=0 else None
    totals.update(net_collections=totals['collections']-totals['refunds'] if has_data else None,attendance_rate=round(totals['attendance']/totals['inquiry']*100,2) if has_data else None,
        cost_per_linked_attendance=round(totals['net_spend']/totals['attendance']) if totals['attendance'] and totals['expense_entries'] and not incomplete and totals['net_spend']>=0 else None,roas=None,complete_cases=complete,incomplete_cases=incomplete)
    warnings=['연결키별 문의 건의 관측입니다. 고유 환자수·병원 신규환자수·광고 인과효과와 다릅니다.','수납·환불은 업로드 기록이며 EMR 금액 대사와 광고 귀속 검증 전이므로 ROAS를 계산하지 않습니다.','집행비는 문의기간의 장부 지출·환불이며 수납은 문의일부터 선택한 관측기간 안의 기록입니다.']
    if incomplete:warnings.append(f'{incomplete}건은 {window}일 관측이 끝나지 않았습니다. 관측된 건수만 표시하며 최종 성과로 보지 않습니다.')
    if orphan:warnings.append(f'문의 기록에 연결되지 않은 {orphan}개 연결키는 전환율·수납 합산에서 제외했습니다.')
    if not has_data:warnings.append('선택기간에 시작한 문의 연결자료가 없습니다. 관측이 없는 값은 성과 0으로 확정하지 않습니다.')
    suggestions=[]
    if orphan:suggestions.append({'reason':f'문의 미연결 {orphan}건','action':'동일한 무작위 연결키로 문의 이벤트가 제공됐는지 확인하세요.'})
    if incomplete:suggestions.append({'reason':f'관측 미완료 {incomplete}건','action':'선택한 관측기간이 지난 뒤 기준일을 갱신한 자료로 다시 확인하세요.'})
    for row in rows.values():
        if row['net_spend']>0 and not row['inquiry']:suggestions.append({'reason':row['title']+' · 집행비 있으나 문의 연결자료 없음','action':'성과를 판단하기 전에 해당 캠페인의 연결자료 누락을 확인하세요.'})
        if row['inquiry']>=5 and row['booking']<row['inquiry']:suggestions.append({'reason':f'{row["title"]} · 문의 {row["inquiry"]}건, 예약 확인 {row["booking"]}건','action':'예약 기록 누락과 실제 미예약 사유를 구분해 다음 응대 실험을 정하세요.'})
    return {'revision':revision,'workspace_revision':workspace['revision'],'checked_at':now(),'has_data':has_data,'period':{'start':start,'end':end},'window_days':window,
        'observed_through':min(cutoffs) if cutoffs else None,'latest_event':max((x['occurred_at'][:10] for x in events),default=None),'complete_cases':complete,'incomplete_cases':incomplete,'orphan_cases':orphan,
        'cohort_event_first':min(observed_dates,default=None),'cohort_event_last':max(observed_dates,default=None),'observation_end_min':min(window_ends,default=None),'observation_end_max':max(window_ends,default=None),
        'totals':totals,'campaigns':sorted(rows.values(),key=lambda x:(-x['inquiry'],x['title'])),'warnings':warnings,'suggestions':suggestions,'stages':STAGES}


def template_xlsx():
    book=openpyxl.Workbook();sheet=book.active;sheet.title='연결자료';sheet.append(list(HEADERS))
    for column in 'ABCDEFG':sheet.column_dimensions[column].width=24
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()
