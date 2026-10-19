"""Read-only Excel snapshots with revalidation on every dashboard request."""
import hashlib
import json
import math
import re
import threading
from collections import Counter
from dataclasses import dataclass, replace
from datetime import date, datetime, timezone, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from .regions import RegionIndex
from .road_addresses import ROAD_DB

KST=timezone(timedelta(hours=9))
MONEY_COLUMNS={
    'treatment_total':'총진료비','insurer':'공단부담','copay':'본인부담',
    'noncovered':'비급여','tax':'부가세','discount':'할인액',
    'card':'카드수납','cash':'현금수납','other':'기타(온라인)',
}


class DataError(Exception):
    """Safe, value-free message suitable for the local UI."""


def norm(value):return re.sub(r'\s+','',str(value if value is not None else ''))


def patient_key(value):
    if value is None or not str(value).strip():return None
    if isinstance(value,(int,float)) and math.isfinite(value) and float(value).is_integer():return str(int(value))
    return str(value).strip()


def parse_date(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    m=re.search(r'((?:19|20)\d{2})[-./년 ]+(\d{1,2})[-./월 ]+(\d{1,2})',str(value or ''))
    if not m:return None
    try:return date(*map(int,m.groups()))
    except ValueError:return None


def parse_money(value):
    if value is None or str(value).strip() in ('','-'):return None
    text=str(value).replace(',','').replace('₩','').strip()
    multiplier=1
    for suffix,factor in [('억원',100000000),('만원',10000),('천원',1000),('원',1)]:
        if text.endswith(suffix):
            text=text[:-len(suffix)].strip();multiplier=factor;break
    try:
        amount=Decimal(text)
        if not amount.is_finite():raise InvalidOperation
        return amount*multiplier
    except InvalidOperation:raise DataError('금액으로 읽을 수 없는 셀이 있습니다. 파일 형식을 확인해 주세요.') from None


def channel(value):
    text=norm(value).lower()
    if not text:return '미기록'
    for keywords,label in [
        (('cpa',),'CPA'),(('강남언니',),'강남언니'),(('여신티켓',),'여신티켓'),
        (('네이버','naver'),'네이버'),(('인스타','instagram'),'인스타그램'),
        (('유튜브','youtube'),'유튜브'),(('블로그',),'블로그'),(('인터넷','검색','홈페이지'),'인터넷·검색'),
        (('지인','소개','가족'),'소개·지인'),(('간판','지나','워킹','방문'),'간판·방문'),
        (('카카오',),'카카오'),(('전단','현수막','버스','옥외'),'오프라인 광고'),
    ]:
        if any(k in text for k in keywords):return label
    return '기타 경로'


def age_sex(value):
    text=str(value or '')
    match=re.search(r'\d{1,3}',text)
    age=int(match.group()) if match else None
    if age is not None and not 0<=age<=120:age=None
    band='미기록' if age is None else ('80세 이상' if age>=80 else f'{age//10*10}대')
    sex='여성' if any(x in text for x in ('여','女','女')) else '남성' if any(x in text for x in ('남','男')) else '미기록'
    return band,sex


def table(sheet):
    iterator=sheet.iter_rows(values_only=True)
    headers=[norm(x) for x in next(iterator,())]
    for index,values in enumerate(iterator,2):
        if any(v is not None for v in values):yield index,dict(zip(headers,values))


@dataclass(frozen=True)
class Snapshot:
    clinic: str
    records: list
    patients: dict
    reports: dict
    inventory: list
    quality: dict
    index: RegionIndex
    checked_at: str
    loaded_at: str
    fingerprint: str
    address_source: dict


class WorkbookStore:
    def __init__(self,manifest):
        self.manifest=Path(manifest)
        self._lock=threading.RLock()
        self._cache=None
        self.index=RegionIndex(road_path=ROAD_DB)
        self.imports=None

    def enable_imports(self,root):
        from .imports import ImportManager
        self.imports=ImportManager(self,root)

    def _configuration(self):
        config=json.loads(self.manifest.read_text(encoding='utf-8'))
        overrides=self.imports.overrides() if self.imports else {}
        config['files']=list(dict.fromkeys(config['files']+list(overrides)))
        return config,overrides

    def import_baseline(self):
        """An empty baseline is allowed only for newly created upload-only clinics."""
        with self._lock:
            try:
                config,overrides=self._configuration()
                if config.get('allow_empty') is True and not config['files'] and not overrides:
                    fingerprint=hashlib.sha256(json.dumps(config,sort_keys=True).encode()).hexdigest()
                    checked=datetime.now(KST).isoformat()
                    return Snapshot(str(config['clinic']),[],{}, {},[],{},self.index,checked,checked,fingerprint,{})
                return self.snapshot()
            except (OSError,ValueError,KeyError,TypeError):raise DataError('업로드 기준 자료를 확인하지 못했습니다.') from None

    def snapshot(self):
        with self._lock:
            try:
                config,overrides=self._configuration()
                source=Path(config['source_directory']).resolve()
                names=config['files']
                if not names and config.get('allow_empty') is True:raise DataError('접수수납목록 엑셀을 먼저 업로드해 주세요. 데이터 업데이트에서 시작할 수 있습니다.')
                if not names or len(names)!=len(set(names)):raise DataError('파일 목록이 비어 있거나 중복돼 있습니다.')
                inventory=[]; paths=[]
                for name in names:
                    if Path(name).name!=name or not name.endswith('.xlsx'):raise DataError('허용되지 않은 파일 이름입니다.')
                    path=overrides.get(name,(source/name).resolve())
                    if (name not in overrides and path.parent!=source) or not path.is_file():raise DataError(f'원본 파일을 찾을 수 없습니다: {name}')
                    stat=path.stat();digest=hashlib.sha256(path.read_bytes()).hexdigest()
                    inventory.append({'file':name,'bytes':stat.st_size,'modified_at':datetime.fromtimestamp(stat.st_mtime,KST).isoformat(),'sha256':digest,'status':'확인'})
                    paths.append(path)
                address_fingerprint=self.index.refresh()
                fingerprint=hashlib.sha256(json.dumps([config.get('clinic'),address_fingerprint,[(x['file'],x['sha256'],x['modified_at']) for x in inventory]]).encode()).hexdigest()
                checked=datetime.now(KST).isoformat()
                if self._cache and self._cache.fingerprint==fingerprint:
                    metadata={item['file']:item for item in self._cache.inventory}
                    inventory=[{**metadata[item['file']],**item} for item in inventory]
                    return replace(self._cache,checked_at=checked,inventory=inventory,address_source=dict(self.index.address_source))
                snapshot=self._load(config,paths,inventory,fingerprint,checked)
                # Reject a mixed snapshot if a source changed during parsing.
                if any(hashlib.sha256(p.read_bytes()).hexdigest()!=item['sha256'] for p,item in zip(paths,inventory)):
                    raise DataError('읽는 중 파일이 변경됐습니다. 저장을 마친 뒤 새로고침해 주세요.')
                if self.index.refresh()!=address_fingerprint:
                    raise DataError('읽는 중 공개 주소 DB가 변경됐습니다. 갱신을 마친 뒤 새로고침해 주세요.')
                self._cache=snapshot
                return snapshot
            except DataError:raise
            except (OSError,ValueError,KeyError,TypeError):
                raise DataError('원본 파일 또는 설정을 읽지 못했습니다. 파일 상태를 확인해 주세요.') from None

    def _load(self,config,paths,inventory,fingerprint,checked):
        records=[];patients={};reports={};quality={'duplicate_rows':0,'errors_in_workbooks':0,'first_visit_missing':0,'money_missing_cells':0,'money_negative_cells':0,'address_conflicts':0}
        for path,item in zip(paths,inventory):
            name=path.name
            year_match=re.match(r'(20\d{2})',name)
            year=int(year_match.group(1)) if year_match else None
            try:book=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
            except Exception:raise DataError(f'엑셀을 열지 못했습니다: {name}') from None
            item['sheets']=len(book.worksheets);item['rows']=sum(s.max_row or 0 for s in book.worksheets)
            try:
                if re.fullmatch(r'20\d{2} 접수수납목록.xlsx',name):
                    seen=set()
                    for _,r in table(book.worksheets[0]):
                        required={'접수시각','차트번호','신구','최초내원'}
                        if not required.issubset(r):raise DataError(f'접수목록 열 구조가 다릅니다: {name}')
                        day=parse_date(r['접수시각']);key=patient_key(r['차트번호'])
                        if day is None or key is None:raise DataError(f'접수일 또는 차트번호가 없는 행이 있습니다: {name}')
                        if day.year!=year:raise DataError(f'파일 연도와 접수일이 다릅니다: {name}')
                        signature=tuple(str(v) for v in r.values())
                        if signature in seen:quality['duplicate_rows']+=1;continue
                        seen.add(signature)
                        amounts={k:parse_money(r.get(col)) for k,col in MONEY_COLUMNS.items()}
                        quality['money_missing_cells']+=sum(v is None for v in amounts.values())
                        quality['money_negative_cells']+=sum(v is not None and v<0 for v in amounts.values())
                        first=parse_date(r['최초내원']);quality['first_visit_missing']+=int(first is None)
                        band,sex=age_sex(r.get('연령성별'))
                        records.append({'key':key,'date':day,'year':year,'new':norm(r['신구'])=='신환','first':first,'channel':channel(r.get('내원경로')),'age':band,'sex':sex,'amounts':amounts})
                elif '내원환자내원경로분포' in name:
                    summaries=list(table(book.worksheets[0]))
                    totals=next((r for _,r in summaries if norm(r.get('내원경로'))=='합계'),{})
                    reports.setdefault(year,{})['channel_totals']={k:float(v) if v is not None else None for k in ('내원환자수','신환수','총내원횟수','총진료비') for v in [parse_money(totals.get(k))]}
                    for sheet in book.worksheets[1:]:
                        for _,r in table(sheet):
                            key=patient_key(r.get('차트번호'))
                            if key is None:continue
                            address=str(r.get('주소') or '').strip()
                            previous=patients.get(key)
                            if previous and previous['address'] and address and previous['address']!=address:quality['address_conflicts']+=1
                            if previous and previous['source_year']>year:continue
                            if previous and previous['address'] and not address:continue
                            patients[key]={'address':address,'source_year':year,'region':self.index.match(address)}
                elif any(token in name for token in ('월별접수환자수','월별예약현황','내원환자연령분포','상담자별상담통계','수가별통계','진료항목별통계')):
                    reports.setdefault(year,{})[self._report_kind(name)]=self._report(name,book.worksheets[0])
                else:
                    errors=0
                    for sheet in book.worksheets:
                        for row in sheet.iter_rows():errors+=sum(c.data_type=='e' for c in row)
                    item['error_cells']=errors;quality['errors_in_workbooks']+=errors
                    item['role']='참고 자료 · 접수목록과 중복 합산 안 함'
            finally:book.close()
        if not records:raise DataError('접수수납목록에 읽을 수 있는 기록이 없습니다.')
        records.sort(key=lambda r:r['date'])
        for r in records:
            p=patients.get(r['key'])
            r['region']=p['region'] if p else {'sig':None,'emd':None,'address_present':False}
        return Snapshot(str(config.get('clinic','병원')),records,patients,reports,inventory,quality,self.index,checked,checked,fingerprint,dict(self.index.address_source))

    @staticmethod
    def _report_kind(name):
        for token,kind in [('월별접수환자수','monthly'),('월별예약현황','reservations'),('내원환자연령분포','age_report'),('상담자별상담통계','consultations'),('수가별통계','fees'),('진료항목별통계','treatments')]:
            if token in name:return kind

    def _report(self,name,sheet):
        kind=self._report_kind(name);result=[]
        fields={
            'monthly':{'days':'진료일수','new_patients':'신환수','patients':'접수환자수','visits':'총접수횟수'},
            'reservations':{'days':'예약일수','general':'일반예약','checkup':'정기검진','scaling':'치석제거','new':'신환예약','fulfillment':'예약이행율(%)'},
            'age_report':{'patients':'내원환자수','existing':'구환수','new':'신환수','average_treatment':'1인당평균진료비'},
            'consultations':{'plans':'계획수','confirmed':'확정건수','rate':'확정비율','patients':'상담환자수','confirmed_patients':'확정환자수','confirmed_amount':'확정금액'},
            'fees':{'patients':'환자수','visits':'진료횟수','entries':'입력횟수','amount':'진료금액'},
            'treatments':{'patients':'환자수','visits':'진료횟수','amount':'진료금액'},
        }[kind]
        for i,(_,r) in enumerate(table(sheet),1):
            first=next(iter(r.values()),None)
            if norm(first) in ('합계','평균','총계',''):continue
            if kind in ('monthly','reservations'):
                m=re.fullmatch(r'(\d{1,2})월?',norm(first))
                if not m:continue
                label=f'{int(m.group(1))}월'
            elif kind=='consultations':label=f'상담자 {i:02d}'
            elif kind=='age_report':
                label=str(first).strip() if re.fullmatch(r'[\d\s~세대이상하미만초과\-]+',str(first)) else '기타 연령'
            else:
                original=str(r.get('수가명') if kind=='fees' else first or '')
                label=next((word for word in ('임플란트','라미네이트','교정','보철','보존','치주','외과','검진','예방','기타') if word in original),None)
                if kind=='fees':label=f'수가 항목 {i:03d}'
                elif label is None:label=f'진료항목 {i:02d}'
            entry={'label':label}
            for k,col in fields.items():
                value=r.get(col)
                if isinstance(value,str):value=value.strip().rstrip('%')
                amount=parse_money(value)
                entry[k]=float(amount) if amount is not None else None
            result.append(entry)
        return result
