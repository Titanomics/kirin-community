"""Local versioned annual Excel imports. Values never leave the aggregate API."""
import hashlib
import io
import json
import os
import re
import sqlite3
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

import openpyxl
from .loader import DataError, KST, MONEY_COLUMNS, norm

MAX_BYTES=25*1024*1024
SCHEMAS={
    'receipts':('접수수납목록',['접수시각','차트번호','신구','최초내원']),
    'channels':('내원환자내원경로분포',['내원경로','내원환자수','신환수','총내원횟수','총진료비']),
    'monthly':('월별접수환자수',['월','신환수','접수환자수','총접수횟수']),
    'reservations':('월별예약현황',['월','일반예약','정기검진','치석제거','신환예약','예약이행율(%)']),
    'age_report':('내원환자연령분포',['내원환자수','구환수','신환수','1인당평균진료비']),
    'consultations':('상담자별상담통계',['계획수','확정건수','확정비율','상담환자수','확정환자수','확정금액']),
    'fees':('수가별통계',['수가명','환자수','진료횟수','입력횟수','진료금액']),
    'treatments':('진료항목별통계',['환자수','진료횟수','진료금액']),
}


def now():return datetime.now(KST).isoformat()
def digest(data):return hashlib.sha256(data).hexdigest()


class ImportManager:
    def __init__(self,store,root):
        self.store=store
        self.root=Path(root).resolve()
        self.state_path=self.root/'state.json'
        self.state_marker=self.root/'state-initialized'
        self.pending={}
        config=json.loads(store.manifest.read_text(encoding='utf-8'))
        self.clinic_id=config.get('clinic_id') or digest(str(store.manifest.resolve()).encode())[:16]

    def _state(self):
        if not self.state_path.exists():
            if self.state_marker.exists():raise DataError('확정된 업로드 이력이 없습니다. 과거 원본으로 되돌아가지 않습니다. 이력을 복구해 주세요.')
            return {'clinic_id':self.clinic_id,'revision':0,'active':{},'history':[]}
        try:
            state=json.loads(self.state_path.read_text(encoding='utf-8'))
            if state['clinic_id']!=self.clinic_id:raise ValueError
            if type(state['revision'])!=int or state['revision']<0 or not isinstance(state['active'],dict) or not isinstance(state['history'],list):raise ValueError
            if not self.state_marker.exists():self.state_marker.write_text('initialized\n',encoding='ascii')
            return state
        except (OSError,ValueError,KeyError):raise DataError('업로드 이력을 읽지 못했습니다. 다른 병원 자료를 합치지 않습니다.') from None

    def _path(self,relative):
        path=(self.root/relative).resolve()
        if not path.is_relative_to(self.root):raise DataError('업로드 저장 위치가 올바르지 않습니다.')
        return path

    def overrides(self):return {name:self._path(info['path']) for name,info in self._state()['active'].items()}

    def _save(self,state):
        self.root.mkdir(parents=True,exist_ok=True)
        temporary=self.root/(uuid.uuid4().hex+'.json.tmp')
        try:
            with temporary.open('w',encoding='utf-8') as out:
                json.dump(state,out,ensure_ascii=False,allow_nan=False);out.flush();os.fsync(out.fileno())
            self.state_marker.write_text('initialized\n',encoding='ascii')
            os.replace(temporary,self.state_path)
        finally:temporary.unlink(missing_ok=True)

    @contextmanager
    def _write_guard(self):
        # SQLite's OS lock serializes independent servers sharing this import directory.
        self.root.mkdir(parents=True,exist_ok=True)
        connection=sqlite3.connect(self.root/'write-lock.sqlite',timeout=25,isolation_level=None)
        try:
            connection.execute('BEGIN IMMEDIATE')
            yield
        finally:
            connection.rollback();connection.close()

    def _open(self,data):
        if not isinstance(data,bytes) or not 0<len(data)<=MAX_BYTES:raise DataError('엑셀 파일은 25MB 이하로 선택해 주세요.')
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries=archive.infolist()
                if len(entries)>5000 or sum(x.file_size for x in entries)>150*1024*1024:raise ValueError
                if any(x.flag_bits&1 for x in entries):raise ValueError
                if any(x.filename.endswith('vbaProject.bin') for x in entries):raise ValueError
                # Excel reports need no DTD/entity expansion or external relationship evaluation.
                for item in entries:
                    if item.filename.endswith(('.xml','.rels')):
                        content=archive.read(item).upper()
                        if b'<!DOCTYPE' in content or b'<!ENTITY' in content:raise ValueError
            book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            if len(book.worksheets)>150:book.close();raise ValueError
            return book
        except Exception:raise DataError('지원하는 .xlsx 파일이 아닙니다. 암호·매크로 없이 다시 내보내 주세요.') from None

    def inspect(self,filename,data):
        book=self._open(data)
        try:
            headers=[norm(x)[:100] for x in next(book.worksheets[0].iter_rows(values_only=True),())]
            if len(headers)>150:raise DataError('열 수가 너무 많습니다.')
            known=set(MONEY_COLUMNS.values())|{'연령성별','내원경로','월','진료일수','예약일수','주소','이름','성별','생년월일','방문일','고객ID','구분','첫방문'}
            known.update(h for _,required in SCHEMAS.values() for h in required)
            columns=[{'id':f'column:{i}','label':h if h in known else f'{i+1}열 (미확인)'} for i,h in enumerate(headers)]
            name=Path(str(filename)).name
            kind=next((key for key,(token,_) in SCHEMAS.items() if token in name),None)
            if not kind:kind=next((key for key,(_,required) in SCHEMAS.items() if set(required).issubset(headers)),None)
            match=re.search(r'20\d{2}',name)
            return {'headers':[column['label'] for column in columns],'columns':columns,'kind':kind,'year':int(match.group()) if match else None,
                    'schemas':{key:{'label':label,'required':required,'optional':list(MONEY_COLUMNS.values())+['연령성별','내원경로'] if key=='receipts' else []} for key,(label,required) in SCHEMAS.items()}}
        finally:book.close()

    def _normalize(self,data,path,kind,mapping):
        book=self._open(data);out=openpyxl.Workbook(write_only=True);row_count=0;saved=False
        try:
            for index,sheet in enumerate(book.worksheets):
                target=out.create_sheet(f'보고서{index+1}')
                rows=sheet.iter_rows(); first=next(rows,())
                source_headers=[norm(cell.value) for cell in first]
                if len(source_headers)>150 or len(set(h for h in source_headers if h))!=len([h for h in source_headers if h]):
                    raise DataError('열 이름이 중복되거나 열 수가 너무 많습니다.')
                if any(cell.data_type in ('f','e') for cell in first):raise DataError('열 제목에 수식·오류값이 있습니다.')
                applicable=mapping if index==0 else {};positions={}
                for key,value in applicable.items():
                    if value is None:continue
                    if value.startswith('column:') and value[7:].isdigit():position=int(value[7:])
                    elif value in source_headers:position=source_headers.index(value)
                    else:raise DataError('매핑한 원본 열을 찾지 못했습니다.')
                    if not 0<=position<len(source_headers) or position in positions:raise DataError('열 연결이 잘못됐거나 중복됐습니다.')
                    positions[position]=key
                excluded={key for key,value in applicable.items() if value is None}
                headers=[positions.get(i,f'제외열{i+1}' if h in excluded or h in applicable else h) for i,h in enumerate(source_headers)]
                if len(set(h for h in headers if h))!=len([h for h in headers if h]):raise DataError('서로 다른 항목을 같은 열로 연결할 수 없습니다.')
                required=SCHEMAS[kind][1] if index==0 else (['차트번호','주소'] if kind=='channels' else [])
                if not set(required).issubset(headers):raise DataError('필수 열이 없습니다: '+', '.join(sorted(set(required)-set(headers))))
                order=list(range(len(headers)))
                if index==0 and kind in ('monthly','reservations'):
                    month_index=headers.index('월');order=[month_index]+[i for i in order if i!=month_index]
                target.append([headers[i] for i in order]);nonempty=0;months=set()
                for row in rows:
                    row_count+=1
                    if row_count>200000:raise DataError('한 번에 처리할 수 있는 행 수를 초과했습니다.')
                    if any(cell.data_type in ('f','e') for cell in row):raise DataError('수식·오류값이 있는 파일입니다. 값으로 내보낸 보고서를 사용해 주세요.')
                    values=[cell.value for cell in row]
                    if not any(v is not None for v in values):continue
                    nonempty+=1
                    if index==0 and kind=='receipts' and norm(values[headers.index('신구')]) not in ('신환','구환'):
                        raise DataError('신구 항목은 신환·구환으로 표시된 보고서가 필요합니다.')
                    month_text=norm(values[order[0]])
                    summary_row=re.fullmatch(r'(?:(?:20\d{2})년?)?(?:합계|총계)|월?평균',month_text)
                    if index==0 and kind in ('monthly','reservations') and month_text and not summary_row:
                        month=re.fullmatch(r'(\d{1,2})월?',month_text)
                        if not month or not 1<=int(month.group(1))<=12 or int(month.group(1)) in months:raise DataError('월별 보고서의 월이 잘못됐거나 중복됐습니다.')
                        months.add(int(month.group(1)))
                    target.append([values[i] if i<len(values) else None for i in order])
                if index==0 and (not nonempty or (kind in ('monthly','reservations') and not months)):
                    raise DataError('내용이 없는 보고서로 기존 자료를 교체할 수 없습니다.')
            out.save(path);saved=True
        finally:
            book.close()
            if not saved:
                # Finalize the streaming workbook so temporary XML writers are closed on validation errors.
                try:out.save(io.BytesIO())
                except Exception:pass
            out.close()

    def _candidate(self,active,checked):
        config=json.loads(self.store.manifest.read_text(encoding='utf-8'))
        names=list(dict.fromkeys(config['files']+list(active)))
        if not names and config.get('allow_empty') is True:return None
        paths=[self._path(active[name]['path']) if name in active else Path(config['source_directory'])/name for name in names]
        inventory=[]
        for name,path in zip(names,paths):
            if not path.is_file():raise DataError('연결된 원본 파일이 없습니다. 파일을 복구한 뒤 다시 확인해 주세요.')
            stat=path.stat()
            inventory.append({'file':name,'bytes':stat.st_size,'modified_at':datetime.fromtimestamp(stat.st_mtime,KST).isoformat(),'sha256':digest(path.read_bytes()),'status':'확인'})
        return self.store._load(config,paths,inventory,'candidate',checked)

    def preview(self,filename,data,*,year,kind,mapping=None):
        with self.store._lock:
            if kind not in SCHEMAS or type(year)!=int or not 2000<=year<=datetime.now(KST).year:raise DataError('보고서 종류와 연도를 확인해 주세요.')
            mapping=mapping or {}
            targets=set(SCHEMAS[kind][1])|set(MONEY_COLUMNS.values())|{'연령성별','내원경로'}
            if not isinstance(mapping,dict) or set(mapping)-targets or any(v is not None and not isinstance(v,str) for v in mapping.values()):raise DataError('열 연결을 확인해 주세요.')
            mapped=[v for v in mapping.values() if v is not None]
            if len(set(mapped))!=len(mapped):raise DataError('열 연결이 중복됐습니다.')
            state=self._state();current=self.store.import_baseline();identifier=uuid.uuid4().hex
            token=SCHEMAS[kind][0];name=f'{year} 접수수납목록.xlsx' if kind=='receipts' else f'{year}년{token}.xlsx'
            original=next((item for item in current.inventory if item['file']==name),None)
            active=state['active'].get(name)
            raw_digest=digest(data)
            duplicate=(active and active['source_sha256']==raw_digest and active.get('mapping',{})==mapping) or (not active and original and original['sha256']==raw_digest and not mapping)
            folder=self.root/'versions'/identifier;folder.mkdir(parents=True,exist_ok=False)
            path=folder/name
            try:
                self._normalize(data,path,kind,mapping)
                item={'path':path.relative_to(self.root).as_posix(),'source_sha256':raw_digest,'mapping':mapping,'updated_at':now(),'kind':kind,'year':year}
                proposed={**state['active'],name:item}
                candidate=self._candidate(proposed,now())
                from .analytics import build_dashboard
                checked=build_dashboard(candidate,start=f'{year}-01-01',end=f'{year}-12-31')
                records=[r for r in candidate.records if r['year']==year]
                previous=[r for r in current.records if r['year']==year]
                quality=checked['quality']
                preview={'id':identifier,'revision':state['revision'],'clinic_id':self.clinic_id,
                         'files':[{'name':name,'kind':kind,'year':year,'action':'duplicate' if duplicate else 'replace' if original else 'add',
                                   'records':len(records) if kind=='receipts' else None,'previous_records':len(previous) if kind=='receipts' else None,
                                   'start':min((r['date'] for r in records),default=None).isoformat() if records else None,
                                   'end':max((r['date'] for r in records),default=None).isoformat() if records else None}],
                         'quality':{'monthly_checks':quality['monthly_checks'],'monthly_passed':quality['monthly_passed'],
                                    'monthly_mismatches':quality['monthly_mismatches'],'duplicate_rows':quality['duplicate_rows'],
                                    'financial':quality['financial']},'checked_at':now()}
                self.pending[identifier]={'preview':preview,'active':proposed,'baseline':current.fingerprint,'name':name,'duplicate':bool(duplicate),'path':path,'sha256':digest(path.read_bytes())}
                # Keep pending tokens bounded; expired previews can be prepared again.
                while len(self.pending)>20:self.pending.pop(next(iter(self.pending)))
                return preview
            except Exception:
                path.unlink(missing_ok=True)
                if folder.exists() and not any(folder.iterdir()):folder.rmdir()
                raise

    def commit(self,identifier,revision,*,confirmed=False):
        with self.store._lock,self._write_guard():
            if not confirmed:raise DataError('연도별 전체 보고서의 교체 범위를 확인해 주세요.')
            pending=self.pending.get(identifier);state=self._state()
            if not pending or revision!=state['revision'] or pending['preview']['revision']!=revision:raise DataError('미리보기가 만료됐습니다. 다시 검사해 주세요.')
            if self.store.import_baseline().fingerprint!=pending['baseline']:raise DataError('검사 후 원본이 변경됐습니다. 다시 검사해 주세요.')
            if digest(pending['path'].read_bytes())!=pending['sha256']:raise DataError('검사 후 업로드 파일이 변경됐습니다. 다시 검사해 주세요.')
            if pending['duplicate']:
                self.pending.pop(identifier);return {'status':'duplicate','revision':revision}
            state['history'].append({'at':now(),'action':'import','file':pending['name'],'before':state['active'],'after':pending['active']})
            state['active']=pending['active'];state['revision']+=1
            self._save(state);self.pending.pop(identifier);self.store._cache=None
            return {'status':'committed','revision':state['revision']}

    def rollback(self,revision):
        with self.store._lock,self._write_guard():
            state=self._state()
            if revision!=state['revision'] or not state['history'] or state['history'][-1]['action']!='import':raise DataError('되돌릴 최근 반영 내역이 없거나 이력이 변경됐습니다.')
            previous=state['history'][-1]['before'];self._candidate(previous,now())
            state['history'].append({'at':now(),'action':'rollback','file':state['history'][-1]['file'],'before':state['active'],'after':previous})
            state['active']=previous;state['revision']+=1;self._save(state);self.store._cache=None
            return {'status':'rolled_back','revision':state['revision']}

    def status(self):
        with self.store._lock:
            state=self._state()
            return {'clinic_id':self.clinic_id,'revision':state['revision'],'checked_at':now(),
                    'can_rollback':bool(state['history'] and state['history'][-1]['action']=='import'),
                    'active':[{'file':name,'year':item['year'],'kind':item['kind'],'updated_at':item['updated_at']} for name,item in state['active'].items()],
                    'history':[{k:entry[k] for k in ('at','action','file')} for entry in reversed(state['history'][-30:])]}
